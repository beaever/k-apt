"""4단계: 수집/파싱 결과를 엑셀로 출력 (10개 컬럼 + 상태).

"상태"는 사용자가 어떤 행을 확인해야 하는지 바로 알 수 있도록 3단계로 나눈다.
- 정상: 10개 컬럼이 모두 실제 값이고 값을 의심할 신호가 없음 -> 확인하지 않아도 됨
- 확인 필요: 공고문은 읽었지만 값을 못 찾았거나 오인 가능성이 있음 -> 셀 메모의 근거 문장만 보면 됨
- 수집 실패: 페이지/첨부파일을 못 받았거나 공고문을 읽지 못함 -> 원문 링크로 직접 확인

자본금/실적/공법번호 셀에는 값을 뽑은 원문 문장과 파일명을 셀 메모로 붙인다.
"""

import re

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill

from app.crawler.step2_download import DETAIL_URL_TMPL
from app.parser.step3_parse import (
    CAPITAL_AMOUNT,
    CAPITAL_KEYWORD,
    NEEDS_CHECK_PREFIX,
    NOT_APPLICABLE,
    normalize_amount,
)

COLUMNS = [
    "공고일자", "아파트명", "공사명", "세대수", "자본금",
    "실적", "공법번호", "낙찰방법", "낙찰업체", "낙찰금액", "상태",
]

STATUS_OK = "정상"
STATUS_CHECK = "확인 필요"
STATUS_FAIL = "수집 실패"

# 값 대신 "사유"가 들어갈 수 있는 컬럼 -> parse_files()가 반환하는 필드 키
PARSED_FIELDS = {"자본금": "capital", "실적": "record", "공법번호": "patent_no"}
# 웹 결과 표에 근거를 넘길 때 쓰는 키 접두사 (프런트는 "_"로 시작하는 키를 컬럼으로 보여주지 않음)
EVIDENCE_KEY = "_근거_"

# K-apt 목록의 단지명 뒤에 붙는 발주 주체 표기 (예: "서서울삼성아파트 (입대의)")
APT_SUFFIX = re.compile(r"\s*\(입대의\)\s*$")
RECORD_YEARS = re.compile(r"(?<!\d)(\d{1,2})\s*(?:개\s*)?년")
RECORD_HOUSEHOLDS = re.compile(r"(\d[\d,]*)\s*세대")
RECORD_COUNT = re.compile(r"(\d+)\s*(?:[건회]|개(?!\s*년))")
# 참가자격 요건이 아니라 제출서류 목록("실적증명서 1부")에서 뽑힌 문장일 가능성
DOCUMENT_LINE = re.compile(r"증명서|사본|협약서|\d+\s*부(?!\S)")
AMOUNT_UNITS = {"억": 1, "천만": 0.1, "백만": 0.01, "만": 0.0001}
# 아파트 도로포장 입찰에서 상식적인 요건 범위. 벗어나면 숫자를 잘못 읽었을 가능성이 크다.
# ponytail: 고정 범위 -> 다른 공종(대형 공사 등)으로 넓히면 범위 조정 필요
CAPITAL_RANGE_EOK = (0.1, 50)
RECORD_MAX_COUNT = 30

OK_FONT = Font(color="FF1E7E34", bold=True)
CHECK_FILL = PatternFill(start_color="FFFFF3CD", end_color="FFFFF3CD", fill_type="solid")
CHECK_FONT = Font(color="FF856404", bold=True)
FAIL_FILL = PatternFill(start_color="FFF8D7DA", end_color="FFF8D7DA", fill_type="solid")
FAIL_FONT = Font(color="FFB02A37", bold=True)
LINK_FONT = Font(color="FF0563C1", underline="single")


def summarize_capital(sentence: str) -> str:
    """"4) 자본금 : 5억 이상인 업체" -> "5억". 금액을 못 찾으면 원문 그대로."""
    keyword = CAPITAL_KEYWORD.search(sentence)
    m = CAPITAL_AMOUNT.search(sentence, keyword.end() if keyword else 0) or CAPITAL_AMOUNT.search(sentence)
    return normalize_amount(m.group()) if m else sentence


def summarize_record(sentence: str) -> str:
    """실적 요건 문장 -> "500세대 5건"(세대 조건이 있으면) 또는 "5년 5건". 건수를 못 찾으면 원문 그대로."""
    count = RECORD_COUNT.search(sentence)
    if not count:
        return sentence
    cond = RECORD_HOUSEHOLDS.search(sentence)
    cond = f"{cond.group(1)}세대" if cond else ""
    if not cond and (years := RECORD_YEARS.search(sentence)):
        cond = f"{years.group(1)}년"
    return f"{cond} {re.sub(r'\s+', '', count.group())}".strip()


def _amount_in_eok(amount: str) -> float | None:
    """"5억", "5천만원", "500,000,000원" -> 억 단위 숫자. 해석 못 하면 None."""
    m = re.fullmatch(r"([\d,.]+)(억|천만|백만|만)?원?", amount)
    if not m:
        return None
    try:
        number = float(m.group(1).replace(",", ""))
    except ValueError:
        return None
    return number * AMOUNT_UNITS[m.group(2)] if m.group(2) else number / 1e8


def _suspicions(r: dict, capital: str, record: str) -> list[str]:
    """값은 찾았지만 오인 가능성이 있는 신호들."""
    reasons = []
    amounts = r.get("capital_amounts") or []
    if len(amounts) > 1:
        reasons.append(f"공고문에 자본금 금액이 여러 개({', '.join(amounts)})")
    eok = _amount_in_eok(capital)
    if eok is not None and not CAPITAL_RANGE_EOK[0] <= eok <= CAPITAL_RANGE_EOK[1]:
        reasons.append(f"자본금 {capital}은 일반적인 범위를 벗어남")
    count = RECORD_COUNT.search(record)
    if count and not 1 <= int(count.group(1)) <= RECORD_MAX_COUNT:
        reasons.append(f"실적 {count.group()}은 일반적인 범위를 벗어남")
    for col, key in (("자본금", "capital"), ("실적", "record")):
        sentence = r.get(key, "")
        if r.get(f"{key}_found") and DOCUMENT_LINE.search(sentence) and "이상" not in sentence:
            reasons.append(f"{col}이 제출서류 문장에서 뽑혔을 수 있음")
    return reasons


def _evidence(r: dict, key: str) -> str:
    """셀 메모로 보여줄 근거: 값을 뽑은 원문 문장 + 파일명."""
    sentence = r.get("patent_no_evidence") if key == "patent_no" else r.get(key, "")
    sentence = (sentence or "").removeprefix(NEEDS_CHECK_PREFIX)
    source = r.get(f"{key}_source", "")
    if not source or not sentence or sentence == NOT_APPLICABLE:
        return ""
    return f"근거: {sentence}\n파일: {source}"


def to_row_dict(r: dict) -> dict:
    """수집 결과 1건을 엑셀 컬럼과 동일한 순서의 dict로 변환.

    엑셀 생성과 API 응답(프런트 결과 테이블)이 같은 컬럼 매핑을 쓰도록 공유.
    """
    # 값을 확정적으로 찾은 경우만 요약. "확인 필요"/사유 문구는 사람이 봐야 하므로 원문 유지.
    capital, record = r.get("capital", ""), r.get("record", "")
    if r.get("capital_found") and capital != NOT_APPLICABLE:
        capital = summarize_capital(capital)
    if r.get("record_found") and record != NOT_APPLICABLE:
        record = summarize_record(record)

    values = {
        "공고일자": r.get("announce_date", "")[:10],
        "아파트명": APT_SUFFIX.sub("", r.get("apt_name") or r.get("apt_name_list", "")),
        "공사명": r.get("work_name") or r.get("bid_title", ""),
        "세대수": r.get("household_count", ""),
        "자본금": capital,
        "실적": record,
        "공법번호": r.get("patent_no", ""),
        "낙찰방법": r.get("bid_method", ""),
        "낙찰업체": r.get("winner", ""),
        "낙찰금액": r.get("bid_amount") or r.get("winner_amount", ""),
    }

    if not r.get("success", True):
        # 브라우저 오류는 뒤에 여러 줄짜리 내부 로그가 붙어 첫 줄만 보여준다
        status = f"{STATUS_FAIL}: {(r.get('error') or '알 수 없는 오류').splitlines()[0]}"
    elif r.get("extract_failed"):
        status = f"{STATUS_FAIL}: {r.get('capital', '공고문을 읽지 못함')}"
    else:
        reasons = []
        for col, value in values.items():
            key = PARSED_FIELDS.get(col)
            if key is None:
                if not value:
                    reasons.append(f"{col} 정보 없음")
            elif not r.get(f"{key}_found", bool(value)):
                reasons.append(f"{col}({value or '값 없음'})")
        reasons += _suspicions(r, capital, record)
        if r.get("extract_errors"):
            reasons.append("일부 첨부파일을 읽지 못함(" + "; ".join(r["extract_errors"]) + ")")
        status = f"{STATUS_CHECK}: " + " / ".join(reasons) if reasons else STATUS_OK

    bid_num = r.get("bid_num")
    values["공고 링크"] = DETAIL_URL_TMPL.format(bid_num=bid_num) if bid_num else ""
    values["상태"] = status
    for col, key in PARSED_FIELDS.items():
        if evidence := _evidence(r, key):
            values[EVIDENCE_KEY + col] = evidence
    return values


def build_excel(row_dicts: list[dict], out_path: str) -> None:
    """to_row_dict()로 이미 변환된 행(dict)들을 받아 엑셀로 저장."""
    wb = Workbook()
    ws = wb.active
    ws.title = "낙찰결과"
    ws.append(COLUMNS)
    status_col = COLUMNS.index("상태") + 1
    for row_dict in row_dicts:
        ws.append([row_dict[c] for c in COLUMNS])
        row_no = ws.max_row
        for col in PARSED_FIELDS:
            if evidence := row_dict.get(EVIDENCE_KEY + col):
                cell = ws.cell(row=row_no, column=COLUMNS.index(col) + 1)
                cell.comment = Comment(evidence, "수집기", width=400, height=120)

        status_cell = ws.cell(row=row_no, column=status_col)
        status = row_dict["상태"]
        if status == STATUS_OK:
            status_cell.font = OK_FONT
            continue
        fill, font = (FAIL_FILL, FAIL_FONT) if status.startswith(STATUS_FAIL) else (CHECK_FILL, CHECK_FONT)
        for cell in ws[row_no]:
            cell.fill = fill
        status_cell.font = font
        # 확인이 필요한 행만 아파트명에 K-apt 원문 링크를 건다
        if row_dict.get("공고 링크"):
            apt_cell = ws.cell(row=row_no, column=COLUMNS.index("아파트명") + 1)
            apt_cell.hyperlink = row_dict["공고 링크"]
            apt_cell.font = LINK_FONT
    ws.freeze_panes = "A2"
    wb.save(out_path)
